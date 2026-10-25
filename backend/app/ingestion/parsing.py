"""Parse an uploaded file into ordered, source-referenced segments (blueprint §7.2).

Every segment keeps its page and its heading path (`heading → subsection → body → page`), which later
become the chunk's section breadcrumb and the citation's page range.

- PDF: MuPDF's layout dictionary in reading order (`sort=True`). Headings are lines set in a larger or bold
  font than the body text. Lines repeated in the top/bottom margin of most pages (running headers, footers,
  page numbers) are dropped. Ruled tables are extracted row by row. Pages without text are flagged; scanned
  documents need OCR, which is P1.
- DOCX: body paragraphs and tables in document order (Word headers/footers are not part of the body).
  Headings come from the Title/Heading N styles. Page numbers are estimated from page breaks, so DOCX
  citations rely mainly on the section path.
- TXT: UTF-8 text; form feeds separate pages; Markdown `#` headings are recognized.
"""

import io
import re
import unicodedata
from collections import Counter
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from typing import Any, Literal

import docx
import pymupdf
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.ingestion.errors import ErrorCode, IngestionError
from app.ingestion.normalize import normalize_text
from app.ingestion.redaction import redact
from app.ingestion.validation import SourceType, decode_text

pymupdf.no_recommend_layout()  # otherwise MuPDF prints an install hint to stdout on first use

SegmentKind = Literal["heading", "paragraph", "table"]

MARGIN_FRACTION = 0.08  # top/bottom band where running headers and footers live
REPEAT_FRACTION = 0.5  # a margin line on at least half the pages is a running header/footer
HEADING_SIZE_RATIO = 1.15
MAX_HEADING_CHARS = 120
LOW_CONFIDENCE_RATIO = 0.05  # share of unreadable characters on a page that triggers a warning
_PAGE_NUMBER = re.compile(r"^(?:page\s*)?\d{1,4}(?:\s*(?:/|of)\s*\d{1,4})?$", re.IGNORECASE)
_MD_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*$")


@dataclass(frozen=True, slots=True)
class Segment:
    kind: SegmentKind
    text: str
    page: int  # 1-based
    section_path: tuple[str, ...]
    level: int = 0  # heading level: 0 = document title, 1 = top heading; always 0 for body


@dataclass(slots=True)
class ParsedDocument:
    source_type: SourceType
    segments: list[Segment]
    page_count: int
    warnings: list[dict[str, Any]] = field(default_factory=list)


class _Outline:
    """Tracks the current heading path while walking a document top to bottom."""

    def __init__(self) -> None:
        self._stack: list[tuple[int, str]] = []

    def enter(self, level: int, title: str) -> tuple[str, ...]:
        while self._stack and self._stack[-1][0] >= level:
            self._stack.pop()
        self._stack.append((level, title))
        return self.path

    @property
    def path(self) -> tuple[str, ...]:
        return tuple(title for _, title in self._stack)


def parse_document(data: bytes, source_type: SourceType) -> ParsedDocument:
    parsers = {"pdf": _parse_pdf, "docx": _parse_docx, "txt": _parse_txt}
    parsed = parsers[source_type](data)
    _redact_segments(parsed)
    if not any(s.kind != "heading" for s in parsed.segments):
        raise IngestionError(ErrorCode.FILE_NO_TEXT, "no body text extracted")
    return parsed


def _redact_segments(parsed: ParsedDocument) -> None:
    total: Counter[str] = Counter()
    cleaned = []
    for seg in parsed.segments:
        text, found = redact(seg.text)
        total += found
        path = tuple(redact(title)[0] for title in seg.section_path)
        cleaned.append(Segment(seg.kind, text, seg.page, path, seg.level))
    parsed.segments = cleaned
    if total:
        parsed.warnings.append({"code": "REDACTED", "counts": dict(sorted(total.items()))})


# ---------- PDF ----------


@dataclass(slots=True)
class _Line:
    text: str
    size: float
    bold: bool
    y0: float
    y1: float
    block: int


def _parse_pdf(data: bytes) -> ParsedDocument:
    try:
        pdf = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise IngestionError(ErrorCode.FILE_CORRUPT, type(exc).__name__) from exc
    with pdf:
        pages = [_pdf_page(page) for page in pdf.pages()]

    body_size = _body_font_size(line for lines, _, _ in pages for line in lines)
    repeated = _repeated_margin_lines(pages)
    heading_sizes = sorted(
        {
            round(line.size, 1)
            for lines, _, _ in pages
            for line in lines
            if line.size >= body_size * HEADING_SIZE_RATIO and len(line.text) <= MAX_HEADING_CHARS
        },
        reverse=True,
    )

    outline = _Outline()
    segments: list[Segment] = []
    empty_pages: list[int] = []
    low_confidence: list[int] = []
    for number, (lines, tables, height) in enumerate(pages, start=1):
        kept = [ln for ln in lines if not _is_margin_noise(ln, height, repeated)]
        if not kept and not tables:
            empty_pages.append(number)
            continue
        if _unreadable_ratio("".join(ln.text for ln in kept)) > LOW_CONFIDENCE_RATIO:
            low_confidence.append(number)

        pending_tables = sorted(tables, key=lambda t: t[1])

        def flush_tables(
            before_y: float,
            page: int = number,
            pending_tables: list[tuple[list[str], float]] = pending_tables,
        ) -> None:
            # Tables are placed back into reading order by their vertical position.
            while pending_tables and pending_tables[0][1] <= before_y:
                rows, _ = pending_tables.pop(0)
                text = normalize_text("\n".join(rows), keep_lines=True)
                if text:
                    segments.append(Segment("table", text, page, outline.path))

        for block_lines in _group_blocks(kept):
            flush_tables(block_lines[0].y0)
            level = _heading_level(block_lines, body_size, heading_sizes)
            raw = "\n".join(ln.text for ln in block_lines)
            if level:
                title = normalize_text(raw)
                segments.append(Segment("heading", title, number, outline.enter(level, title), level))
            else:
                text = normalize_text(raw)
                if text:
                    segments.append(Segment("paragraph", text, number, outline.path))
        flush_tables(float("inf"))

    warnings: list[dict[str, Any]] = []
    if empty_pages and len(empty_pages) < len(pages):
        warnings.append({"code": "PAGES_WITHOUT_TEXT", "pages": empty_pages})
    if low_confidence:
        warnings.append({"code": "PAGES_LOW_CONFIDENCE", "pages": low_confidence})
    if len(empty_pages) == len(pages):
        raise IngestionError(ErrorCode.FILE_NO_TEXT, "every page is empty (scanned?)")
    return ParsedDocument("pdf", segments, len(pages), warnings)


def _pdf_page(page: pymupdf.Page) -> tuple[list[_Line], list[tuple[list[str], float]], float]:
    """Text lines outside tables, table rows, and the page height."""
    tables: list[tuple[list[str], float]] = []
    table_boxes: list[pymupdf.Rect] = []
    try:
        for table in page.find_tables().tables:
            # A cell's wrapped lines are one value: "\n" would start a new row in the text (same as DOCX).
            rows = [" | ".join(" ".join((cell or "").split()) for cell in row) for row in table.extract()]
            rows = [r for r in rows if r.strip(" |")]
            if len(rows) >= 2:
                tables.append((rows, table.bbox[1]))
                table_boxes.append(pymupdf.Rect(table.bbox))
    except Exception:  # noqa: S110 - table detection is best effort; text extraction below still runs
        pass

    lines: list[_Line] = []
    layout = page.get_text("dict", sort=True)
    for block_no, block in enumerate(layout["blocks"]):
        if block.get("type") != 0:
            continue
        for line in block["lines"]:
            spans = [s for s in line["spans"] if s["text"].strip()]
            if not spans:
                continue
            box = pymupdf.Rect(line["bbox"])
            if any(box.intersects(t) for t in table_boxes):
                continue
            chars = sum(len(s["text"]) for s in spans)
            size = sum(s["size"] * len(s["text"]) for s in spans) / chars
            bold = all(s["flags"] & 16 or "bold" in s["font"].lower() for s in spans)
            text = "".join(s["text"] for s in spans).strip()
            lines.append(_Line(text, size, bold, box.y0, box.y1, block_no))
    return lines, tables, page.rect.height


def _body_font_size(lines: Iterable[_Line]) -> float:
    sizes: Counter[float] = Counter()
    for line in lines:
        sizes[round(line.size, 1)] += len(line.text)
    return sizes.most_common(1)[0][0] if sizes else 10.0


def _margin_key(text: str) -> str:
    return re.sub(r"\d+", "#", text.strip().lower())


def _in_margin(line: _Line, height: float) -> bool:
    return line.y1 <= height * MARGIN_FRACTION or line.y0 >= height * (1 - MARGIN_FRACTION)


def _repeated_margin_lines(pages: list[tuple[list[_Line], Any, float]]) -> set[str]:
    if len(pages) < 3:
        return set()
    seen: Counter[str] = Counter()
    for lines, _, height in pages:
        seen.update({_margin_key(ln.text) for ln in lines if _in_margin(ln, height)})
    threshold = max(3, int(len(pages) * REPEAT_FRACTION))
    return {key for key, count in seen.items() if count >= threshold}


def _is_margin_noise(line: _Line, height: float, repeated: set[str]) -> bool:
    if not _in_margin(line, height):
        return False
    return _margin_key(line.text) in repeated or bool(_PAGE_NUMBER.fullmatch(line.text.strip()))


def _group_blocks(lines: list[_Line]) -> Iterator[list[_Line]]:
    """Consecutive lines of one layout block, split where the font role changes (heading ↔ body)."""
    group: list[_Line] = []
    for line in lines:
        if group and (line.block != group[-1].block or abs(line.size - group[-1].size) > 0.5):
            yield group
            group = []
        group.append(line)
    if group:
        yield group


def _heading_level(block: list[_Line], body_size: float, heading_sizes: list[float]) -> int:
    text = " ".join(ln.text for ln in block)
    if len(block) > 3 or len(text) > MAX_HEADING_CHARS or text.endswith((".", ",", ";")):
        return 0
    size = round(block[0].size, 1)
    if size in heading_sizes:
        return heading_sizes.index(size) + 1
    if all(ln.bold for ln in block) and size >= body_size - 0.5 and len(block) == 1:
        return len(heading_sizes) + 1  # bold run-in heading at body size: the lowest level
    return 0


def _unreadable_ratio(text: str) -> float:
    if not text:
        return 0.0
    bad = sum(1 for ch in text if ch == "\ufffd" or unicodedata.category(ch) in {"Co", "Cn", "Cc"})
    return bad / len(text)


# ---------- DOCX ----------


def _parse_docx(data: bytes) -> ParsedDocument:
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise IngestionError(ErrorCode.FILE_CORRUPT, type(exc).__name__) from exc

    body = document.element.body
    # Word's own pagination (rendered breaks) is the best estimate; generated files only have hard breaks.
    use_rendered = bool(body.xpath(".//w:lastRenderedPageBreak"))

    outline = _Outline()
    segments: list[Segment] = []
    page = 1
    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            breaks_before, breaks_inside = _docx_page_breaks(item, use_rendered)
            page += breaks_before
            text = normalize_text(item.text)
            if text:
                level = _docx_heading_level(item)
                if level is not None:
                    segments.append(Segment("heading", text, page, outline.enter(level, text), level))
                else:
                    style = (item.style.name if item.style is not None else "") or ""
                    if style.lower().startswith("list") and not text.startswith("- "):
                        text = f"- {text}"
                    segments.append(Segment("paragraph", text, page, outline.path))
            page += breaks_inside
        elif isinstance(item, Table):
            rows = _docx_table_rows(item)
            text = normalize_text("\n".join(rows), keep_lines=True)
            if text:
                segments.append(Segment("table", text, page, outline.path))
    return ParsedDocument("docx", segments, page)


def _docx_heading_level(paragraph: Paragraph) -> int | None:
    """0 for the document title (above every heading), N for Heading N, None for body text."""
    name = (paragraph.style.name if paragraph.style is not None else "") or ""
    if name == "Title":
        return 0
    match = re.fullmatch(r"Heading (\d)", name)
    return int(match.group(1)) if match else None


def _docx_page_breaks(paragraph: Paragraph, use_rendered: bool) -> tuple[int, int]:
    """(breaks before the paragraph's text, breaks after it)."""
    p = paragraph._p
    before = 1 if p.xpath("./w:pPr/w:pageBreakBefore[not(@w:val='0')]") else 0
    if use_rendered:
        return before + len(paragraph.rendered_page_breaks), 0
    hard = [br for br in p.xpath(".//w:br") if br.get(qn("w:type")) == "page"]
    return before, len(hard)


def _docx_table_rows(table: Table) -> list[str]:
    rows: list[str] = []
    for row in table.rows:
        cells: list[str] = []
        previous = None
        for cell in row.cells:
            if cell._tc is previous:  # horizontally merged cells repeat the same element
                continue
            previous = cell._tc
            cells.append(" ".join(cell.text.split()))
        if any(cells):
            rows.append(" | ".join(cells))
    return rows


# ---------- TXT ----------


def _parse_txt(data: bytes) -> ParsedDocument:
    text = decode_text(data)
    outline = _Outline()
    segments: list[Segment] = []
    pages = text.split("\f")
    for number, page_text in enumerate(pages, start=1):
        for block in re.split(r"\n\s*\n", page_text.replace("\r\n", "\n")):
            lines = block.strip("\n").split("\n")
            heading = _MD_HEADING.fullmatch(lines[0].strip()) if lines and lines[0].strip() else None
            if heading:
                title = normalize_text(heading.group(2))
                level = len(heading.group(1))
                segments.append(Segment("heading", title, number, outline.enter(level, title), level))
                lines = lines[1:]
            body = normalize_text("\n".join(lines))
            if body:
                segments.append(Segment("paragraph", body, number, outline.path))
    return ParsedDocument("txt", segments, len(pages))
