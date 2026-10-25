"""Build small PDF, DOCX and TXT files in memory, so tests never depend on binary fixtures in the repo."""

import io
from collections.abc import Sequence

import docx
import pymupdf
from docx.enum.text import WD_BREAK

PageSpec = Sequence[tuple[str, str]]  # (kind, text): kind is "h1", "h2", "p" or "table" ("a|b\nc|d")


def pdf_bytes(
    pages: Sequence[PageSpec],
    *,
    header: str | None = None,
    footer_page_numbers: bool = False,
    password: str | None = None,
) -> bytes:
    """A text PDF. Headings are set larger than body text; tables are drawn as ruled grids."""
    doc = pymupdf.open()
    for number, spec in enumerate(pages, start=1):
        page = doc.new_page(width=595, height=842)
        if header:
            page.insert_text((50, 30), header, fontsize=8)
        if footer_page_numbers:
            page.insert_text((290, 825), f"Page {number} of {len(pages)}", fontsize=8)
        y = 80.0
        for kind, text in spec:
            if kind == "table":
                y = _draw_table(page, y, [row.split("|") for row in text.split("\n")])
                continue
            size = {"h1": 20, "h2": 15}.get(kind, 10.5)
            box = pymupdf.Rect(50, y, 545, 800)
            # insert_textbox returns the unused height; measure what was used to place the next block.
            spare = page.insert_textbox(box, text, fontsize=size, fontname="helv")
            used = (box.height - spare) if spare >= 0 else box.height
            y += used + size * 0.9
    out: bytes = doc.tobytes(
        encryption=pymupdf.PDF_ENCRYPT_AES_256 if password else pymupdf.PDF_ENCRYPT_KEEP,  # type: ignore[attr-defined]
        owner_pw=password,
        user_pw=password,
    )
    doc.close()
    return out


TABLE_FONT_SIZE = 9
TABLE_LINE_HEIGHT = 11


def _draw_table(page: pymupdf.Page, y: float, rows: list[list[str]]) -> float:
    """Ruled grid; long cell text wraps inside its cell and the row grows to fit, as in a real document (an
    unwrapped line ran past the cell edge, and table extraction rightly kept only the part inside it)."""
    columns = max(len(r) for r in rows)
    width = 495 / columns
    for row in rows:
        cells = [_wrap(row[c].strip() if c < len(row) else "", width - 6) for c in range(columns)]
        height = 7 + TABLE_LINE_HEIGHT * max(len(lines) for lines in cells)
        for c, lines in enumerate(cells):
            cell = pymupdf.Rect(50 + c * width, y, 50 + (c + 1) * width, y + height)
            page.draw_rect(cell, color=(0, 0, 0), width=0.6)
            for i, line in enumerate(lines):
                page.insert_text((cell.x0 + 3, cell.y0 + 2 + TABLE_LINE_HEIGHT * (i + 1)), line, fontsize=9)
        y += height
    return y + 14


def _wrap(text: str, width: float) -> list[str]:
    lines: list[str] = []
    for word in text.split():
        candidate = f"{lines[-1]} {word}" if lines else word
        if lines and pymupdf.get_text_length(candidate, fontsize=TABLE_FONT_SIZE) <= width:
            lines[-1] = candidate
        else:
            lines.append(word)
    return lines or [""]


def scanned_pdf_bytes(pages: int = 2) -> bytes:
    """A PDF whose pages hold only an image: what a scanner produces without OCR."""
    doc = pymupdf.open()
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 40), False)
    pix.clear_with(200)
    for _ in range(pages):
        page = doc.new_page()
        page.insert_image(pymupdf.Rect(50, 50, 300, 300), pixmap=pix)
    out: bytes = doc.tobytes()
    doc.close()
    return out


def docx_bytes(blocks: Sequence[tuple[str, str]]) -> bytes:
    """kinds: "title", "h1", "h2", "p", "list", "table" ("a|b\\nc|d"), "pagebreak"."""
    document = docx.Document()
    for kind, text in blocks:
        if kind == "title":
            document.add_heading(text, level=0)
        elif kind in ("h1", "h2", "h3"):
            document.add_heading(text, level=int(kind[1]))
        elif kind == "list":
            document.add_paragraph(text, style="List Bullet")
        elif kind == "table":
            rows = [r.split("|") for r in text.split("\n")]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = cell.strip()
        elif kind == "pagebreak":
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        else:
            document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def warranty_pdf() -> bytes:
    """A fictional three-page TechCare policy with headings, a table, a running header and page numbers."""
    return pdf_bytes(
        [
            [
                ("h1", "TechCare Laptop Warranty"),
                (
                    "p",
                    "This warranty covers TechCare laptops bought from TechCare stores or techcare.example.",
                ),
                ("h2", "Coverage period"),
                (
                    "p",
                    "Laptops are covered for 12 months from the date of purchase. "
                    "Batteries are covered for 6 months.",
                ),
            ],
            [
                ("h2", "What is not covered"),
                (
                    "p",
                    "Accidental damage, liquid damage and cosmetic wear are not covered "
                    "by the standard warranty.",
                ),
                ("table", "Plan|Length|Accidental damage\nStandard|12 months|No\nCare+|24 months|Yes"),
            ],
            [
                ("h2", "How to claim"),
                (
                    "p",
                    "Bring the laptop and your receipt to any TechCare store. "
                    "Claims take up to 10 business days.",
                ),
            ],
        ],
        header="TechCare Laptop Warranty - internal copy",
        footer_page_numbers=True,
    )
